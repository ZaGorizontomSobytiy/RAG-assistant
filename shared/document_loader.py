"""
Загрузка документов из файлов и директорий.
Поддерживаемые форматы: .txt, .md, .pdf, .docx.
"""

import os
import re
from pathlib import Path
from typing import List, Optional
from dataclasses import dataclass

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
ENCODING = "utf-8"


@dataclass
class Document:
    """Один документ с метаданными для индексации."""
    text: str
    source: str
    filename: str
    mtime: float = 0.0


def _read_text_file(path: Path) -> str:
    with open(path, "r", encoding=ENCODING) as f:
        return f.read()


def _load_txt_md(path: Path) -> Optional[Document]:
    try:
        text = _read_text_file(path)
        text = text.strip()
        if not text:
            return None
        return Document(
            text=text,
            source=str(path.resolve()),
            filename=path.name,
            mtime=path.stat().st_mtime,
        )
    except Exception as e:
        print(f"  [пропуск] {path}: {e}")
        return None


def _load_pdf(path: Path) -> Optional[Document]:
    try:
        import fitz
        doc = fitz.open(path)
        parts = []
        for page in doc:
            parts.append(page.get_text())
        doc.close()
        text = "\n\n".join(p.strip() for p in parts if p.strip())
        if not text:
            return None
        return Document(
            text=text,
            source=str(path.resolve()),
            filename=path.name,
            mtime=path.stat().st_mtime,
        )
    except Exception as e:
        print(f"  [пропуск] {path}: {e}")
        return None


def _load_docx(path: Path) -> Optional[Document]:
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(parts).strip()
        if not text:
            return None
        return Document(
            text=text,
            source=str(path.resolve()),
            filename=path.name,
            mtime=path.stat().st_mtime,
        )
    except Exception as e:
        print(f"  [пропуск] {path}: {e}")
        return None


def _load_file(path: Path) -> Optional[Document]:
    suf = path.suffix.lower()
    if suf in (".txt", ".md"):
        return _load_txt_md(path)
    if suf == ".pdf":
        return _load_pdf(path)
    if suf == ".docx":
        return _load_docx(path)
    return None


def load_from_path(path: str, recursive: bool = True) -> List[Document]:
    """
    Загрузить документы из файла или директории.

    Args:
        path: путь к файлу или папке
        recursive: при обходе директории заходить в подпапки

    Returns:
        список документов с text, source, filename, mtime
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Путь не найден: {path}")

    documents: List[Document] = []

    if p.is_file():
        if p.suffix.lower() in SUPPORTED_EXTENSIONS:
            doc = _load_file(p)
            if doc:
                documents.append(doc)
        else:
            print(f"  [пропуск] неподдерживаемый формат: {p.name}")
        return documents

    # Директория
    pattern = "**/*" if recursive else "*"
    for f in p.glob(pattern):
        if not f.is_file() or f.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        doc = _load_file(f)
        if doc:
            documents.append(doc)

    return documents
